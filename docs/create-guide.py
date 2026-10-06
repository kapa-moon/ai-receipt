from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json

root=Path(__file__).resolve().parents[1]
doc=Document(); sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(7);doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(27);doc.styles['Heading 1'].font.size=Pt(19);doc.styles['Heading 2'].font.size=Pt(13)
doc.core_properties.title='AI receipt calculations';doc.core_properties.author='AI Receipt project';doc.core_properties.subject='Usage accounting and reflection scoring'
def p(text,style=None): return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,level=1)
def sub(text):doc.add_heading(text,level=2)
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cs=t.add_row().cells
  for i,x in enumerate(row):cs[i].text=str(x)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders);marg=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');marg.append(e)
   pr.append(marg)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'333333' if ri==0 else ('F4F4F4' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1.05
    if ci>0 and headers[ci] not in ['Calculation','Formula or meaning']:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in par.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0);r.bold=ri==0
  if ri==0:
   trpr=row._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 p('')
 return t

doc.add_heading('AI receipt calculations',0)
p('Usage accounting and reflection scoring',style='Subtitle')
p('October 6 2026')
p('This guide explains every numeric measure in the AI Receipt interface: recorded tokens, API-equivalent dollar estimates, messages, tools, category shares, model shares, reflection scores and subscription quota observations. It is intended for people trying the prototype and for research review.')
p('Token and tool counts come from supported trace records. Dollar values are API-price scenarios, not subscription bills. Reflection scores come from users’ answers and are illustrative measures of perceived progress, not earnings, financial ROI or validated economic utility.')
h('Which records are counted')
p('First choose a vendor, then a chat and a date scope. All dates is the default. The date filter selects complete message turns by the user-message timestamp in the browser’s local timezone; it does not slice individual model responses out of a selected turn. Last 24 hours, 3 days, 7 days and 30 days are trailing periods ending at the current time.')
p('All chats combines that vendor’s owner-facing topics. Automatic-review activity remains available as a separate topic and is excluded from All chats. The shared demo is explicitly curated: 12 owner-selected messages and their linked usage are omitted, leaving 10 Lite AI messages. The remaining owner messages are unchanged. New uploaded traces are imported in full; this demo policy is not applied to visitors’ imports.')
sub('Counts in the receipt header')
table(['Measure','Calculation'],[['User requests','Count user-message entries in the selected scope. Background entries without a recorded user prompt do not increase this count.'],['Model responses','Count unique metered response records linked to selected turns. A response can include text, reasoning and tools.'],['Tool activity','Count unique recorded tool calls or nested invocation expressions. This is not proof that the tools succeeded.'],['Files and entries','Files counts source artifacts, including curated receipt exports. Entries counts all shown turns, including background entries.']], [1.45,5.05])
p('Source coverage dates describe the retained receipt records. Refresh time is when the data was normalized. Deleted demo entries are removed before all visible counts, shares and costs are calculated.')

doc.add_page_break();h('How tokens are counted')
sub('OpenAI rollout records')
p('The adapter reads token_usage_record.payload.usage once per response_id. Input includes any cached input. Fresh input = input_tokens − cached_input_tokens − cache_write_input_tokens, with a lower limit of zero. Total recorded tokens use the response total_tokens when present, otherwise input_tokens + output_tokens.')
p('Reasoning tokens are already part of output tokens. They are never added again. Cumulative turn, thread and token_count counters are not added to response totals. A repeated context can be counted again on each request because the provider processes it again.')
sub('Claude transcript and audit records')
p('The adapter reads assistant message usage. Fresh input, cache reads and cache creation are disjoint. Recorded total = input_tokens + cache_read_input_tokens + cache_creation_input_tokens + output_tokens. Thinking tokens are a subset of output, not additional tokens.')
p('Streamed rows with the same request ID or message ID are counted once. The adapter retains the final or largest reported output count and merges unique tool IDs. Cache creation is split by the recorded 5-minute and 1-hour duration fields; the remainder is treated as duration unknown.')
sub('A worked token example')
table(['Recorded quantity','Tokens'],[['Input including cached input','1,200,000'],['Cached input','900,000'],['Fresh input after subtracting cache','300,000'],['Output including reasoning','100,000'],['Reasoning subset already inside output','25,000'],['Total recorded tokens','1,300,000']], [4.8,1.7])
p('The total is 1,300,000, not 1,325,000. Cache reads remain part of the total because cached context is still processed; they usually have a lower price.')
sub('Missing or duplicated data')
p('Invalid or incomplete usage records are skipped and counted in source notes. Duplicate records are discarded and counted separately. A message without supported usage shows missing metering, not proof of free work. Reported zero reasoning does not prove that the model performed no internal reasoning.')
p('The supplied OpenAI traces have zero cache-write tokens. The adapter’s nonzero OpenAI cache-write interpretation would need schema verification before relying on another runtime format.')

doc.add_page_break();h('How API estimates are calculated')
p('For each priced response, cost = (fresh input × input rate + cache reads × cache-read rate + 5-minute writes × 5-minute rate + 1-hour writes × 1-hour rate + output × output rate) / 1,000,000. Each response uses its recorded model. The rate card is a current-rate scenario checked during prototype development, not a historical invoice.')
table(['Model','Input','Cache read','Output'],[['gpt-6-astra','$10.00','$1.00','$50.00'],['gpt-6.1-sol','$2.00','$0.10','$10.00'],['claude-opus-5-5','$4.00','$0.20','$20.00'],['claude-opus-4-8','$5.00','$0.50','$25.00'],['claude-fable-5','$10.00','$1.00','$50.00'],['claude-sonnet-5','$2.00','$0.20','$10.00'],['claude-sonnet-5-5','$2.00','$0.20','$10.00']],[3.2,1.1,1.1,1.1])
p('Rates above are USD per million tokens, as implemented in trace-engine.js. The rates and model identifiers are assumptions to verify before using this prototype for billing decisions.')
sub('Why some estimates have a range')
p('OpenAI lower scenarios use the table rates. The long-context scenario uses input/read/output rates of 20/2/75 for gpt-6-astra and 4/0.2/15 for gpt-6.1-sol. The importer does not resolve the billing tier. These two scenarios are not guaranteed lower and upper bounds on a real bill; speed, contracts, region and other charges can differ.')
p('For Claude, 5-minute cache writes use 1.25 × input price, and 1-hour writes use 2 × input price. Unknown cache duration creates a range between those prices. Recorded fast mode applies a 2 × multiplier; recorded US inference applies a 1.1 × multiplier. Missing speed defaults to normal as an assumption.')
p('An explicitly recorded Claude server web search adds $0.01 per request. Client commands, file edits and browser actions receive no invented flat fees. OpenAI wrapper search calls do not reliably establish billable API search requests, so unverified search and compute charges are omitted.')
sub('Worked estimate and totals')
p('Using the preceding token example with gpt-6.1-sol: lower cost = (300,000 × 2 + 900,000 × 0.1 + 100,000 × 10) / 1,000,000 = $1.69. The long-context scenario is $2.88. A total sums the unrounded costs of all selected responses separately for each scenario.')
p('Unknown models remain unpriced. If every response is unpriced, the interface says Unpriced. Mixed totals show the known subtotal plus an unpriced remainder. Dollars display two decimal places; a small nonzero cost can display $0.00. Internal sums retain full precision.')

doc.add_page_break();h('Category model and action numbers')
sub('Use cases')
p('The default category is inferred from local keyword rules over the beginning of the prompt. It is not an OpenAI proprietary classifier and does not cost API tokens. The optional mini-model tool can supply a category and activity summary; its additional input and output tokens are recorded separately and excluded from the original work receipt.')
table(['Number','Formula or meaning'],[['Token share and bar width','Category recorded tokens / all selected recorded tokens × 100. Bars use the full total, not the largest category.'],['Cost share','Category lower-scenario estimate / selected lower-scenario estimate × 100. Hidden when unknown pricing prevents a complete share.'],['Messages','User messages assigned to that category. Background turns can contribute usage without increasing message count.'],['Cost per message','Category estimate / category user-message count, separately for lower and higher scenarios. No messages means no per-message average.'],['Category total cost','Sum of unrounded API costs for that category’s responses. Categories are mutually exclusive.']], [1.5,5.0])
p('Token and dollar shares differ when a category uses a more expensive model, more output, or less cached context. At zero total tokens the bar is empty. Cost shares use only the lower pricing scenario consistently across the table.')
sub('Models')
p('Model segment width = that model’s selected recorded tokens / all selected recorded tokens × 100. The hover tooltip shows the model’s unique metered responses, recorded tokens and summed API estimate. Model response count can differ from user-message count because one message may generate many model responses.')
sub('AI actions')
p('The Total number row sums unique tool calls; associated token and cost columns are not summed because they overlap. An action count is the number of unique calls grouped into search/retrieve, write/edit files, run commands, browser/interface, ask/clarify, plan/coordinate, discover tools, present results or other tools. Codex exec expressions identify attempted nested calls; a failed or conditional invocation is not proof of success.')
p('In the expanded accounting table, an action’s associated tokens and costs come from responses linked to any of its tool IDs. Each response is counted once within that action. The same response can belong to several actions, so action columns overlap and must not be added. Unlinked calls show Unlinked usage rather than invented tokens.')


doc.add_page_break();h('Reflection scoring without dollar estimates')
p('The interface asks three questions: was the result useful, what moved forward, and what happened to the AI output. We use them to calculate a reflection gain score. The word gain refers to reported usefulness and progress; it does not mean financial profit or causal business improvement.')
sub('Per message usefulness')
p('Usefulness percentage = star rating / 5 × 100. One star is 20%, four stars 80%, and five stars 100%. An unrated message is missing data, not zero. Clicking the selected star again clears the rating.')
sub('Per message reflection gain')
p('Reflection gain = 100 × (0.50 × U + 0.20 × P + 0.30 × O), where U = stars / 5, P is the progress-category weight and O is the output-status weight. All three answers must be present; otherwise the score is not calculated.')
table(['What moved forward','P'],[['Decision','1.0'],['Working artifact','1.0'],['Community connection','1.0'],['Understanding','0.8'],['Errand or operation','0.6'],['Others','0.5'],['Nothing moved forward','0.0']],[4.8,1.7])
table(['What happened to the output','O'],[['Used or applied','1.0'],['Ready to use','0.8'],['Needs revision','0.4'],['Still exploring','0.2'],['Lost or abandoned','0.0']],[4.8,1.7])
p('These category and status weights are transparent prototype choices, not empirical findings or proof that decisions are intrinsically more valuable than errands. They should be validated or revised with users. Under the current complete-response options, possible scores run from 10% to 100%. An abandoned output can still score above zero when the user reports useful learning or progress.')
p('Example: 4 stars, Decision, Still exploring gives 100 × (0.5 × 0.8 + 0.2 × 1 + 0.3 × 0.2) = 66%. Changing only the output status to Used or applied raises it to 90%. No dollar amount is required.')


doc.add_page_break();h('Totals quota and interpretation')
sub('The three totals at the bottom')
p('Total estimate cost sums selected response costs before display rounding. Total estimate utility is the average usefulness percentage over rated user messages only. Reflection gain is the arithmetic mean of complete message reflection scores. Neither reflection total is weighted by token count or money spent. Background activity is excluded from reflection denominators.')
p('The labels show coverage, such as 2 of 3 messages rated and 2 of 3 messages fully reflected. Incomplete reflections do not silently contribute zeros. A high score based on one message does not describe all unassessed work.')
p('Example: message A has 4 stars, Decision and Still exploring, so its gain is 66%. Message B has 5 stars, Understanding and Used or applied, so its gain is 96%. Message C is unassessed. Total utility = (80 + 100) / 2 = 90%; reflection gain = (66 + 96) / 2 = 81%. Both denominators are two, with coverage shown as 2 of 3.')
sub('Subscription quota is a different measure')
p('Quota used percentages come directly from provider-reported observations or manual entries. Remaining percentage = max(0, min(100, 100 − used percentage)). A 300-minute window is displayed as 5 hours; 10,080 minutes is displayed as 7 days. Window duration and reset timestamps are provider fields, not reconstructed from dollars or tokens.')
p('A reset time is shown in the viewer’s local timezone. A stale reset time does not automatically turn a saved observation into zero usage. Charts display sparse observations, with breaks at reset changes, decreases or gaps longer than 30 minutes. The 5-hour, 24-hour and 7-day views select observations relative to the latest recorded sample; they do not alter the receipt’s message scope.')
p('Account-verified and identity-unverified observations remain separate. These account-wide percentages cannot be attributed entirely to Lite AI or assumed to cover every ordinary ChatGPT chat. A missing quota window is not zero usage.')
sub('What these numbers cannot establish')
p('The receipt does not measure revenue, actual time saved, correction time, subscriptions allocated to a project, infrastructure expenses or a causal counterfactual. Financial ROI would require independently estimated benefits and complete costs. The reflection score should be interpreted as a conversation aid for assessing work, not an economic conclusion or a validated Schumacher measure.')
sub('Implementation and references')
p('Code modules: trace-engine.js, receipt-view.js, trace-ui.js, quota-engine.js and curate-data.js. Shared-demo exclusions are listed in demo-policy.json.')
for label,url in [('OpenAI usage insights and assessing value','https://learn.chatgpt.com/docs/enterprise/usage-insights#assessing-value'),('OpenAI API pricing','https://developers.openai.com/api/docs/pricing'),('Claude API pricing','https://platform.claude.com/docs/en/about-claude/pricing'),('Claude Code monitoring','https://code.claude.com/docs/en/monitoring-usage')]:p(label+'\n'+url)
for par in doc.paragraphs[-4:]:
 par.paragraph_format.space_after=Pt(4)
 for run in par.runs:run.font.size=Pt(10)
for tree in [doc.element,doc.styles.element]:
 for border in tree.xpath('.//w:pBdr'):border.getparent().remove(border)
doc.save(root/'docs/AI-receipt-calculations.docx')
